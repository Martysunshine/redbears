"""Text extraction for uploaded artifacts.

Each extractor returns plain text (or raises ``ExtractionError``). Binary files
masquerading as text are rejected rather than silently producing garbage.
"""

from __future__ import annotations

from pathlib import Path

import config


class ExtractionError(Exception):
    """Raised when a file cannot be turned into usable text."""


def kind_for(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in config.PDF_EXTENSIONS:
        return "pdf"
    if ext in config.DOCX_EXTENSIONS:
        return "docx"
    if ext in config.TEXT_EXTENSIONS:
        return "text"
    return "unsupported"


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in config.ARTIFACT_ALLOWED_EXTENSIONS


# ---------------------------------------------------------------------------
# Per-type extractors
# ---------------------------------------------------------------------------

def _extract_pdf(path: Path) -> str:
    import pypdf

    try:
        reader = pypdf.PdfReader(str(path))
    except Exception as exc:  # corrupt / encrypted
        # Messages here reach the uploader verbatim (services/artifacts/service.py
        # stores them on the record), so the library's own words stay on the
        # exception chain for the log and never in the sentence.
        raise ExtractionError(
            "Could not read this PDF. If it is scanned or encrypted, try a text export."
        ) from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(pages).strip()

    # Scanned / image-only PDF: the text layer is empty or negligible, so the
    # content lives only in page pixels. Fall back to OCR (poppler rasterize +
    # Tesseract). Degrades gracefully to the (empty) text layer if OCR is
    # unavailable — extract() then raises the honest "no extractable text".
    if len(text) < config.PDF_OCR_MIN_TEXT_CHARS:
        from services.media import ocr
        ocr_text = ocr.ocr_pdf(path)
        if len(ocr_text) > len(text):
            return ocr_text
    return text


def _extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ExtractionError("DOCX files cannot be processed on this server right now.") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(
            "Could not read this DOCX. Try re-saving it in Word, or upload a text export."
        ) from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _looks_binary(sample: bytes) -> bool:
    if b"\x00" in sample:
        return True
    if not sample:
        return False
    # ratio of non-text control bytes
    text_chars = bytes(range(0x20, 0x7F)) + b"\n\r\t\f\b"
    nontext = sum(1 for b in sample if b not in text_chars)
    return (nontext / len(sample)) > 0.30


def _extract_text(path: Path) -> str:
    raw = path.read_bytes()
    if _looks_binary(raw[:4096]):
        raise ExtractionError("File appears to be binary, not text")
    # try utf-8, then chardet-detected encoding, then latin-1 fallback
    try:
        return raw.decode("utf-8").strip()
    except UnicodeDecodeError:
        pass
    try:
        import chardet

        guess = chardet.detect(raw[:100_000])
        enc = guess.get("encoding")
        if enc and guess.get("confidence", 0) >= 0.5:
            return raw.decode(enc, errors="replace").strip()
    except Exception:
        pass
    return raw.decode("latin-1", errors="replace").strip()


def extract(path: Path, filename: str) -> str:
    """Extract text from a supported file. Raises ExtractionError on failure."""
    kind = kind_for(filename)
    if kind == "pdf":
        text = _extract_pdf(path)
    elif kind == "docx":
        text = _extract_docx(path)
    elif kind == "text":
        text = _extract_text(path)
    else:
        raise ExtractionError(f"Unsupported file type: {Path(filename).suffix}")

    if not text.strip():
        raise ExtractionError("No extractable text found in file")
    # cap very large extractions so embedding/summaries stay bounded
    if len(text) > config.MAX_ARTIFACT_TEXT_CHARS:
        text = text[: config.MAX_ARTIFACT_TEXT_CHARS]
    return text
